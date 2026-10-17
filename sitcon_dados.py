"""
SITCON - camada de dados
Extrai as planilhas brutas, valida com regras explícitas, aplica correções
rastreáveis e devolve: base bruta, base corrigida e catálogo de erros.

Princípios:
- Nada é apagado: cada registro guarda o valor bruto e a célula de origem.
- Formato evidente é corrigido automaticamente (confiança Alta).
- Erro com evidência no próprio dado (vizinhos, padrão) é corrigido por
  inferência (confiança Média) e sinalizado.
- Valor ausente ou ambíguo NUNCA é inventado: fica "Pendente de validação".
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

PASTA = Path(__file__).parent
ARQ_CDVS = PASTA / "levantamento BD CDVs via 1 e 2.xlsx"
ARQ_SINAIS = PASTA / "levantamento BD Sinais via 1 e 2.xlsx"

FUNCOES_CDV = ["NCTP", "SCTP", "N12-CTP", "TPR", "TPRN"]
FUNCOES_SINAL = ["RGE", "RGENX", "YGE", "YGENX"]
PINO_CDV = {"NCTP": "21", "SCTP": "23", "N12-CTP": "26", "TPR": "30", "TPRN": "31"}

RE_CAMPO = re.compile(r"^(LB|LC) ([A-Z]{1,2})-(\d+)$")
RE_TB = re.compile(r"^TB([2-4])-([A-E])(\d+)$")
RE_INTERFACE = re.compile(r"^B(\d{2}) (\d{2})$")
RE_BAC = re.compile(r"^BAC\d+$")

# ---------------------------------------------------------------------------
# Catálogo de regras
# ---------------------------------------------------------------------------
REGRAS = {
    "E01": ("Nome de equipamento duplicado", "Estrutura", "Alta",
            "O mesmo nome aparece em blocos diferentes, com ligações diferentes."),
    "E02": ("Linha duplicada", "Estrutura", "Média",
            "A mesma locação aparece duas vezes no mesmo bloco, com os mesmos valores."),
    "E03": ("Célula com mais de um valor", "Estrutura", "Média",
            "Dois pontos de ligação na mesma célula (ex.: 'LB A-23 / LB B-22')."),
    "E04": ("Rótulo de nível inconsistente", "Padronização", "Baixa",
            "O mesmo nível escrito de formas diferentes (ex.: 'Armario TB', 'Armários TB', 'interface 2')."),
    "E05": ("Código de função irregular", "Padronização", "Média",
            "Código sem o número do equipamento, com espaços ou sufixo divergente."),
    "E06": ("Espaçamento irregular", "Formato", "Baixa",
            "Espaço faltando, sobrando ou no lugar errado (ex.: 'B0923', 'B 59 01', 'TB3-D57 ')."),
    "E07": ("Prefixo inválido", "Formato", "Média",
            "Prefixo de borne fora do padrão LB/LC (ex.: 'BL B-11', 'LB TB-AA22')."),
    "E08": ("Zero à esquerda", "Formato", "Baixa",
            "Número de borne com zero à esquerda (ex.: 'LB B-01'), divergente do restante da base."),
    "E09": ("Endereço de armário inválido", "Formato", "Alta",
            "Endereço TB fora do padrão TB<2-4>-<régua A-E><borne> (ex.: 'TB2-29', 'TB34-A3')."),
    "E10": ("Pino de interface incompleto", "Formato", "Média",
            "Endereço de interface sem o número do pino (ex.: 'B24')."),
    "E11": ("Valor ausente", "Ausência", "Alta",
            "Posição obrigatória sem valor (nem '-'), ou linha de armário vazia."),
    "E12": ("Pino de interface fora do padrão", "Consistência", "Alta",
            "Nos CDVs, os pinos são fixos por função (21/23/26/30/31) e no mesmo bastidor."),
    "E13": ("Barramento N12 inconsistente", "Consistência", "Alta",
            "O N12-CTP usa um borne fixo por locação; este registro diverge da maioria."),
    "E14": ("Espelho Locação 17/18 divergente", "Consistência", "Média",
            "As locações 17 e 18 repetem os mesmos bornes; aqui um valor diverge."),
    "E15": ("Sequência TPR/TPRN quebrada", "Consistência", "Média",
            "O TPRN costuma ser o borne seguinte ao TPR (+1 ou +2); aqui não é."),
    "E16": ("Sequência de aspectos quebrada", "Consistência", "Média",
            "Os 4 bornes de um sinal ficam agrupados; um valor está muito distante dos demais."),
    "E17": ("Conflito de ponto físico", "Conflito", "Alta",
            "O mesmo ponto físico está atribuído a equipamentos ou funções diferentes."),
    "E18": ("Função sem nenhuma ligação", "Completude", "Baixa",
            "A função não tem ponto em nenhum nível (Locação, Armário ou Interface)."),
}

ACAO_AUTO = "Corrigido automaticamente"
ACAO_INFER = "Corrigido por inferência"
ACAO_PEND = "Pendente de validação"


# ---------------------------------------------------------------------------
# Utilitários
# ---------------------------------------------------------------------------
def _txt(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    return str(v)


def celula(linha0: int, coluna0: int) -> str:
    return f"{get_column_letter(coluna0 + 1)}{linha0 + 1}"


def nivel_canonico(rotulo: str):
    r = re.sub(r"\s+", " ", rotulo or "").strip().lower()
    if r.startswith("locação") or r.startswith("locacao"):
        return "Locação"
    if r.startswith("arm"):
        return "Armário TB"
    if r.startswith("interface"):
        return "Interface 2"
    if r.startswith("bastidor"):
        return "Bastidor DCC1"
    return rotulo


def funcao_de_codigo(codigo: str, tipo: str):
    c = (codigo or "").strip().upper()
    opcoes = ["N12-CTP", "NCTP", "SCTP", "TPRN", "TPR"] if tipo == "CDV" else ["RGENX", "RGE", "YGENX", "YGE"]
    for f in opcoes:
        if c.endswith(f):
            return f
    return None


# ---------------------------------------------------------------------------
# 1. Extração estrutural (posição a posição, com célula de origem)
# ---------------------------------------------------------------------------
def _blocos(caminho: Path, tipo: str):
    d = pd.read_excel(caminho, header=None, dtype=object)
    R, C = d.shape
    padrao = r"^\s*CDV\s" if tipo == "CDV" else r"^\s*SINAL\s"
    titulos = [(r, c, str(d.iat[r, c]).strip()) for r in range(R) for c in range(C)
               if isinstance(d.iat[r, c], str) and re.match(padrao, d.iat[r, c], re.I)]
    largura = 5 if tipo == "CDV" else 4
    inicio_dados = 4 if tipo == "CDV" else 5
    blocos = []
    for r, c, t in titulos:
        prox = [rr for rr, cc, _ in titulos if cc == c and rr > r]
        fim = min(prox) if prox else R
        codigos = [_txt(d.iat[r + 3, c + k]) for k in range(1, largura + 1)]
        linhas = []
        for rr in range(r + inicio_dados, fim):
            rot = _txt(d.iat[rr, c])
            vals = [_txt(d.iat[rr, c + k]) if c + k < C else None for k in range(1, largura + 1)]
            if (rot is None or not rot.strip()) and all(v is None or not v.strip() for v in vals):
                continue
            linhas.append((rr, rot, vals))
        blocos.append(dict(titulo=t, linha=r, coluna=c, codigos=codigos, linhas=linhas))
    return d, blocos


def carregar_bruto():
    """Base bruta: uma linha por posição (célula) de ligação, sem nenhuma correção."""
    regs = []
    planilhas = {}
    for arq, tipo in [(ARQ_CDVS, "CDV"), (ARQ_SINAIS, "Sinal")]:
        df, blocos = _blocos(arq, tipo)
        planilhas[tipo] = df
        for b in blocos:
            if tipo == "CDV":
                nome = re.sub(r"^\s*CDV\s+|\s+LUZ\s*$", "", b["titulo"], flags=re.I).strip()
            else:
                nome = "Sinal " + re.sub(r"\D", "", b["titulo"])
            bloco_id = f"{tipo}:{celula(b['linha'], b['coluna'])}"
            for rr, rot, vals in b["linhas"]:
                for k, v in enumerate(vals):
                    regs.append(dict(
                        arquivo=arq.name, tipo=tipo, bloco_id=bloco_id,
                        celula_bloco=celula(b["linha"], b["coluna"]),
                        equipamento_bruto=nome, titulo_bruto=b["titulo"],
                        codigo_bruto=b["codigos"][k], posicao=k,
                        nivel_bruto=rot, valor_bruto=v,
                        celula=celula(rr, b["coluna"] + k + 1),
                        linha_excel=rr + 1,
                    ))
    return pd.DataFrame(regs), planilhas


# ---------------------------------------------------------------------------
# 2. Validação + correção
# ---------------------------------------------------------------------------
class Catalogo:
    def __init__(self):
        self.itens = []

    def add(self, regra, reg, valor_corrigido=None, acao=ACAO_PEND, detalhe="", sugestao=None):
        nome, cat, sev, _ = REGRAS[regra]
        self.itens.append(dict(
            regra=regra, nome_regra=nome, categoria=cat, severidade=sev,
            arquivo=reg.get("arquivo"), tipo=reg.get("tipo"), celula=reg.get("celula"),
            equipamento=reg.get("equipamento_bruto"), bloco=reg.get("celula_bloco"),
            nivel=reg.get("nivel_bruto"), funcao=reg.get("funcao"),
            valor_bruto=reg.get("valor_bruto"), valor_corrigido=valor_corrigido,
            acao=acao, detalhe=detalhe, sugestao=sugestao,
        ))

    def df(self):
        d = pd.DataFrame(self.itens)
        if d.empty:
            return d
        d.insert(0, "id_erro", [f"#{i+1:03d}" for i in range(len(d))])
        return d


def _corrigir_formato(reg, cat: Catalogo):
    """Correções de formato ponto a ponto. Retorna valor corrigido (ou None)."""
    bruto = _txt(reg["valor_bruto"])
    if bruto is None:
        return None
    v = bruto
    if v.strip() in {"-", "—", "–"}:
        return "-"
    if v.strip() == "":
        return None
    # espaços nas pontas
    if v != v.strip():
        cat.add("E06", reg, v.strip(), ACAO_AUTO, "Espaço extra no início/fim da célula.")
        v = v.strip()
    nivel = reg["nivel"]
    if nivel == "Interface 2":
        m = re.match(r"^B\s*(\d{2})\s*(\d{2})$", v)
        if m and not RE_INTERFACE.match(v):
            novo = f"B{m.group(1)} {m.group(2)}"
            cat.add("E06", reg, novo, ACAO_AUTO, "Espaçamento irregular no endereço de interface.")
            v = novo
        else:
            m2 = re.match(r"^B\s*(\d)\s*(\d{2})\s*(\d{2})$", v)  # 'B 59 01' já coberto acima; segurança
            if m2:
                pass
    elif nivel == "Locação":
        if v.startswith("BL "):
            novo = "LB " + v[3:]
            cat.add("E07", reg, novo, ACAO_AUTO, "Prefixo 'BL' invertido; o padrão é 'LB'.")
            v = novo
        m = re.match(r"^(LB|LC) TBB?-([A-Z]{1,2})(\d+)$", v)
        if m:
            novo = f"{m.group(1)} {m.group(2)}-{m.group(3)}"
            cat.add("E07", reg, novo, ACAO_INFER,
                    "Prefixo 'TB-'/'TBB-' dentro do borne de campo; normalizado para o padrão '<régua>-<n>'.")
            v = novo
        m = re.match(r"^(LB|LC) ([A-Z]{1,2})-0+(\d+)$", v)
        if m:
            novo = f"{m.group(1)} {m.group(2)}-{m.group(3)}"
            cat.add("E08", reg, novo, ACAO_AUTO, "Zero à esquerda removido.")
            v = novo
    return v


def _inferir_tb(reg, vizinhos, cat: Catalogo):
    v = _txt(reg["valor"])
    m = re.match(r"^TB(\d*)-([A-Z]*)(\d+)$", v or "")
    if not m:
        cat.add("E09", reg, None, ACAO_PEND, "Endereço não reconhecido.")
        return v
    arm, regua, n = m.group(1), m.group(2), int(m.group(3))
    if len(arm) > 1:
        cat.add("E09", reg, None, ACAO_PEND,
                f"Armário '{arm}' não existe (só TB2, TB3, TB4). Ambíguo — validar em campo.",
                sugestao=" ou ".join(f"TB{a}-{regua}{n}" for a in arm if a in "234"))
        return v
    candidatos = []
    for viz in vizinhos:
        mv = RE_TB.match(viz or "")
        if mv and abs(int(mv.group(3)) - n) <= 2:
            candidatos.append(mv)
    if candidatos:
        a = arm or candidatos[0].group(1)
        r = regua or candidatos[0].group(2)
        novo = f"TB{a}-{r}{n}"
        falta = "número do armário" if not arm else "letra da régua"
        cat.add("E09", reg, novo, ACAO_INFER,
                f"Faltava {falta}; inferido pelo borne vizinho na mesma linha ({candidatos[0].group(0)}).")
        return novo
    cat.add("E09", reg, None, ACAO_PEND, "Endereço incompleto e sem vizinho para inferência.")
    return v


def processar():
    bruto, planilhas = carregar_bruto()
    cat = Catalogo()
    df = bruto.copy()
    df["nivel"] = df["nivel_bruto"].map(nivel_canonico)
    df["locacao"] = df.apply(lambda r: re.sub(r"\s+", " ", r["nivel_bruto"]).strip()
                             if r["nivel"] == "Locação" else None, axis=1)
    df["funcao"] = [funcao_de_codigo(c, t) for c, t in zip(df["codigo_bruto"], df["tipo"])]
    # Sinal 30: códigos sem número -> função pela posição
    for i, r in df[df["funcao"].isna()].iterrows():
        lista = FUNCOES_CDV if r["tipo"] == "CDV" else FUNCOES_SINAL
        df.at[i, "funcao"] = lista[r["posicao"]]

    # --- E04 rótulos inconsistentes (uma ocorrência por variação de rótulo e planilha)
    canon = {"Armário TB": "Armários TB", "Interface 2": "Interface 2"}
    lab = df[df["nivel"].isin(canon)].drop_duplicates("linha_excel" if False else ["bloco_id", "linha_excel"])
    lab = lab[lab.apply(lambda r: r["nivel_bruto"] != canon[r["nivel"]], axis=1)]
    for (arq, rot), g in lab.groupby(["arquivo", "nivel_bruto"]):
        r = g.iloc[0].to_dict()
        r["celula"] = re.sub(r"^[A-Z]+", get_column_letter(_col_bloco(r["celula_bloco"])), r["celula"])
        r["valor_bruto"] = rot
        r["funcao"] = None
        qtd = g["bloco_id"].nunique()
        if qtd > 1:
            r["equipamento_bruto"] = f"{qtd} equipamentos"
        onde = f"em {qtd} blocos" if qtd > 1 else "neste bloco"
        cat.add("E04", r, canon[r["nivel"]], ACAO_AUTO,
                f"'{rot}' usado {onde}, enquanto o padrão adotado é '{canon[r['nivel']]}'. Padronizado.")

    # --- E05 códigos de função
    for bid, g in df.groupby("bloco_id"):
        r0 = g.iloc[0].to_dict()
        codigos = g.drop_duplicates("posicao")["codigo_bruto"].tolist()
        num = re.sub(r"\D", "", r0["equipamento_bruto"]) if r0["tipo"] == "Sinal" else r0["equipamento_bruto"]
        r0["celula"] = f"{get_column_letter(_col_bloco(r0['celula_bloco'])+1)}{int(re.sub(r'[A-Z]', '', r0['celula_bloco']))+3}"
        r0["funcao"] = None
        if r0["tipo"] == "Sinal" and not any(str(c).startswith(num) for c in codigos):
            r0["valor_bruto"] = ", ".join(codigos)
            cat.add("E05", r0, ", ".join(f"{num}{c}" for c in codigos), ACAO_AUTO,
                    "Códigos sem o número do sinal; completados pelo título do bloco.")
        if any(c and c != c.strip() for c in codigos):
            r0["valor_bruto"] = ", ".join(repr(c) for c in codigos)
            cat.add("E05", r0, ", ".join(c.strip() for c in codigos), ACAO_AUTO, "Espaço no início dos códigos.")

    # --- correções de formato ponto a ponto
    valores = []
    for i, r in df.iterrows():
        reg = r.to_dict()
        valores.append(_corrigir_formato(reg, cat))
    df["valor"] = valores

    # --- E03 células multivalor -> explode
    novas = []
    for i, r in df.iterrows():
        v = _txt(r["valor"])
        if v and "/" in v:
            partes = [p.strip() for p in v.split("/") if p.strip()]
            cat.add("E03", r.to_dict(), " | ".join(partes), ACAO_AUTO,
                    f"Separado em {len(partes)} registros (mesma função, pontos distintos).")
            for p in partes:
                nr = r.copy(); nr["valor"] = p; novas.append(nr)
        else:
            novas.append(r)
    df = pd.DataFrame(novas).reset_index(drop=True)

    # --- E09 endereços de armário / E10 pino incompleto (inferência por vizinhos)
    for i, r in df.iterrows():
        v = _txt(r["valor"])
        if not v or v == "-":
            continue
        viz = df[(df["bloco_id"] == r["bloco_id"]) & (df["linha_excel"] == r["linha_excel"]) & (df.index != i)]["valor"].tolist()
        if r["nivel"] == "Armário TB" and not RE_TB.match(v):
            reg = r.to_dict()
            df.at[i, "valor"] = _inferir_tb(reg, viz, cat)
        if r["nivel"] == "Interface 2" and re.match(r"^B\d{2}$", v):
            pinos = sorted(int(m.group(2)) for x in viz if (m := RE_INTERFACE.match(x or "")) and x.startswith(v))
            if len(pinos) >= 2 and len(set(b - a for a, b in zip(pinos, pinos[1:]))) == 1:
                passo = pinos[1] - pinos[0]
                novo = f"{v} {pinos[-1] + passo:02d}"
                cat.add("E10", r.to_dict(), novo, ACAO_INFER,
                        f"Pino ausente; inferido pela sequência da linha ({', '.join(map(str, pinos))}).")
                df.at[i, "valor"] = novo
            else:
                cat.add("E10", r.to_dict(), None, ACAO_PEND, "Pino ausente sem sequência para inferir.")

    # --- E11 valores ausentes
    for bid, g in df.groupby("bloco_id"):
        for lin, gl in g.groupby("linha_excel"):
            vazios = gl[gl["valor"].isna()]
            if vazios.empty:
                continue
            r0 = gl.iloc[0]
            if r0["tipo"] == "Sinal":
                if len(vazios) == len(gl):
                    reg = vazios.iloc[0].to_dict()
                    reg["valor_bruto"] = "(linha inteira vazia)"
                    cat.add("E11", reg, None, ACAO_PEND,
                            f"{r0['nivel']} sem nenhum borne para os 4 aspectos.")
                else:
                    for _, rv in vazios.iterrows():
                        sug = None
                        rge = gl[gl["funcao"] == "RGE"]["valor"].dropna()
                        m = RE_TB.match(rge.iloc[0]) if len(rge) else None
                        if m and rv["funcao"] == "RGENX":
                            sug = f"TB{m.group(1)}-{m.group(2)}{int(m.group(3)) - 1}"
                        cat.add("E11", rv.to_dict(), None, ACAO_PEND,
                                "Aspecto sem borne no armário (vizinhos preenchidos). Não preenchido automaticamente.",
                                sugestao=sug and f"{sug} (padrão RGENX = RGE − 1 em 15 de 19 sinais)")
            else:
                for _, rv in vazios.iterrows():
                    cat.add("E11", rv.to_dict(), None, ACAO_PEND,
                            "Célula vazia num bloco que usa '-' para 'não se aplica'.")

    # --- E02 linhas duplicadas
    remover = []
    for bid, g in df.groupby("bloco_id"):
        assinatura = {}
        for lin, gl in g.groupby("linha_excel"):
            chave = (gl.iloc[0]["nivel_bruto"], tuple(gl.sort_values("posicao")["valor"].fillna("∅")))
            if chave in assinatura:
                reg = gl.iloc[0].to_dict(); reg["valor_bruto"] = f"linha {lin} = linha {assinatura[chave]}"
                cat.add("E02", reg, "linha removida", ACAO_AUTO, "Linha idêntica à anterior; mantida uma cópia.")
                remover.extend(gl.index.tolist())
            else:
                assinatura[chave] = lin
    df = df.drop(index=remover).reset_index(drop=True)

    # --- E01 nomes duplicados -> id único pelo bastidor de interface
    df["equipamento"] = df["equipamento_bruto"]
    blocos_por_nome = df.groupby("equipamento_bruto")["bloco_id"].unique()
    for nome, bids in blocos_por_nome.items():
        if len(bids) > 1:
            for bid in bids:
                g = df[df["bloco_id"] == bid]
                bast = g[(g["nivel"] == "Interface 2") & g["valor"].str.match(r"^B\d{2}", na=False)]["valor"].str[:3].mode()
                sufixo = bast.iloc[0] if len(bast) else bid.split(":")[1]
                novo = f"{nome} [{sufixo}]"
                df.loc[df["bloco_id"] == bid, "equipamento"] = novo
                reg = g.iloc[0].to_dict(); reg["celula"] = g.iloc[0]["celula_bloco"]
                reg["valor_bruto"] = g.iloc[0]["titulo_bruto"]
                cat.add("E01", reg, novo, ACAO_PEND,
                        f"'{nome}' aparece em {len(bids)} blocos com ligações e bastidores diferentes; "
                        f"identificado provisoriamente por {sufixo}. Nome real a confirmar.")

    # --- E12 pinos de interface dos CDVs
    for eq, g in df[(df["tipo"] == "CDV") & (df["nivel"] == "Interface 2")].groupby("equipamento"):
        validos = g[g["valor"].str.match(r"^B\d{2} \d{2}$", na=False)]
        if validos.empty:
            continue
        cont = validos["valor"].str[:3].value_counts()
        forte = cont.iloc[0] >= 3
        bast = cont.index[0]
        usados = set(validos["valor"])
        for i, r in validos.iterrows():
            esperado = f"{bast} {PINO_CDV[r['funcao']]}"
            if forte and r["valor"] == esperado:
                continue
            if forte and r["valor"][:3] == bast and esperado not in usados:
                cat.add("E12", r.to_dict(), esperado, ACAO_INFER,
                        f"{r['funcao']} deve usar o pino {PINO_CDV[r['funcao']]}; o valor bruto repete o pino de outra função.")
                df.at[i, "valor"] = esperado
            elif forte:
                cat.add("E12", r.to_dict(), None, ACAO_PEND,
                        f"Fora do bastidor/pino padrão do CDV. Pode ser ligação real — validar.", sugestao=esperado)
            elif not forte:
                lista = ", ".join(validos["valor"])
                pino_ok = r["valor"].endswith(" " + PINO_CDV[r["funcao"]])
                if pino_ok and r["valor"][:3] == validos.iloc[0]["valor"][:3]:
                    continue
                cat.add("E12", r.to_dict(), None, ACAO_PEND,
                        f"O CDV usa bastidores diferentes ({lista}); nos demais CDVs todas as funções ficam no mesmo bastidor.",
                        sugestao=f"{validos.iloc[0]['valor'][:3]} {PINO_CDV[r['funcao']]} (bastidor da 1ª função; segue a sequência dos vizinhos)")

    # --- Validações de consistência (sobre a base já corrigida)
    _validar_consistencia(df, cat)

    erros = cat.df()
    df["status"] = "OK"
    # marca status do registro
    if not erros.empty:
        por_celula = erros.groupby(["arquivo", "celula"])["acao"].agg(lambda s: ACAO_PEND if ACAO_PEND in set(s) else s.iloc[0])
        df["status"] = [por_celula.get((a, c), "OK") for a, c in zip(df["arquivo"], df["celula"])]
    df["aplicavel"] = df["valor"].ne("-")
    df["via"] = df.apply(lambda r: f"Via {r['equipamento_bruto'][0]}" if r["tipo"] == "CDV" else None, axis=1)
    df["sentido"] = df.apply(lambda r: r["equipamento_bruto"][1] if r["tipo"] == "CDV" else None, axis=1)
    df["armario"] = df["valor"].str.extract(r"^TB([2-4])-")[0].radd("TB").where(df["nivel"] == "Armário TB")
    df["regua"] = df["valor"].str.extract(r"^TB[2-4]-([A-E])")[0].where(df["nivel"] == "Armário TB")
    df["bastidor"] = df["valor"].str.extract(r"^(B\d{2}) ")[0].where(df["nivel"] == "Interface 2")
    df["pino"] = df["valor"].str.extract(r"^B\d{2} (\d{2})$")[0].where(df["nivel"] == "Interface 2")
    return dict(bruto=bruto, corrigido=df, erros=erros, planilhas=planilhas)


def _col_bloco(celula_bloco: str) -> int:
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(re.match(r"^[A-Z]+", celula_bloco).group(0))


def _num(v):
    m = re.search(r"(\d+)$", v or "")
    return int(m.group(1)) if m else None


def _regua(v):
    m = re.match(r"^(?:LB|LC) ([A-Z]{1,2})-", v or "") or re.match(r"^TB\d-([A-E])", v or "")
    return m.group(1) if m else None


def _validar_consistencia(df, cat: Catalogo):
    ok = df["valor"].notna() & df["valor"].ne("-")

    # E13 barramento N12 por locação
    n12 = df[ok & (df["funcao"] == "N12-CTP") & (df["nivel"] == "Locação")]
    for loc, g in n12.groupby("locacao"):
        moda = g["valor"].mode()
        if len(moda) != 1 or g["valor"].nunique() == 1:
            continue
        for _, r in g[g["valor"] != moda.iloc[0]].iterrows():
            qtd = (g["valor"] == moda.iloc[0]).sum()
            cat.add("E13", r.to_dict(), None, ACAO_PEND,
                    f"Na {loc}, {qtd} de {len(g)} CDVs usam {moda.iloc[0]} para o N12; este usa {r['valor']}.")

    # E14 espelho Locação 17/18 (exceto N12, que é por locação)
    campo = df[ok & (df["nivel"] == "Locação") & (df["funcao"] != "N12-CTP")]
    for (eq, fn), g in campo.groupby(["equipamento", "funcao"]):
        l17 = g[g["locacao"] == "Locação 17"]["valor"].tolist()
        l18 = g[g["locacao"] == "Locação 18"]["valor"].tolist()
        if l17 and l18 and set(l17) != set(l18):
            r = g[g["locacao"] == "Locação 17"].iloc[0]
            cat.add("E14", r.to_dict(), None, ACAO_PEND,
                    f"Locação 17 = {', '.join(l17)} | Locação 18 = {', '.join(l18)}. Nos demais casos as duas são iguais.")
    # E15 sequência TPR/TPRN (CDV, mesma linha)
    cdv = df[ok & (df["tipo"] == "CDV") & df["funcao"].isin(["TPR", "TPRN"]) & df["nivel"].isin(["Locação", "Armário TB"])]
    for (eq, lin), g in cdv.groupby(["equipamento", "linha_excel"]):
        if set(g["funcao"]) != {"TPR", "TPRN"}:
            continue
        a = g[g["funcao"] == "TPR"].iloc[0]; b = g[g["funcao"] == "TPRN"].iloc[0]
        if a["valor"] == b["valor"]:
            continue  # mesmo borne: tratado como conflito (E17)
        if _regua(a["valor"]) != _regua(b["valor"]) or (_num(b["valor"]) - _num(a["valor"])) not in (1, 2):
            cat.add("E15", b.to_dict(), None, ACAO_PEND,
                    f"TPR = {a['valor']} e TPRN = {b['valor']}; o padrão da base é TPRN = TPR + 1 ou + 2 na mesma régua.")

    # E16 sequência de aspectos dos sinais (bornes de campo agrupados)
    sin = df[ok & (df["tipo"] == "Sinal") & (df["nivel"] == "Locação")]
    for (eq, lin), g in sin.groupby(["equipamento", "linha_excel"]):
        nums = g["valor"].map(_num)
        med = nums.median()
        desloc = {"RGE": 0, "RGENX": -2, "YGE": 2, "YGENX": 4}
        rge = g[g["funcao"] == "RGE"]["valor"]
        for i, r in g.iterrows():
            if abs(_num(r["valor"]) - med) > 8:
                sug = None
                if len(rge):
                    base = rge.iloc[0]
                    sug = re.sub(r"\d+$", str(_num(base) + desloc[r["funcao"]]), base)
                cat.add("E16", r.to_dict(), None, ACAO_PEND,
                        f"Bornes da linha: {', '.join(g['valor'])}. {r['valor']} está fora do grupo.",
                        sugestao=sug and f"{sug} (padrão RGE=n, RGENX=n−2, YGE=n+2, YGENX=n+4)")

    # E17 conflitos de ponto físico
    pts = df[ok].copy()
    pts["ponto"] = pts.apply(lambda r: f"{r['locacao']} · {r['valor']}" if r["nivel"] == "Locação" else r["valor"], axis=1)
    for p, g in pts.groupby("ponto"):
        donos = g[["equipamento", "funcao"]].drop_duplicates()
        if len(donos) <= 1:
            continue
        if set(donos["funcao"]) == {"N12-CTP"}:
            continue  # barramento comum: compartilhamento por projeto
        nao_n12 = donos[donos["funcao"] != "N12-CTP"]
        n_n12 = (donos["funcao"] == "N12-CTP").sum()
        desc = "; ".join(f"{e} {f}" for e, f in nao_n12.itertuples(index=False))
        extra = f" Colide com o barramento N12 desta locação (usado por {n_n12} CDVs)." if n_n12 else ""
        alvo = g[g["funcao"] != "N12-CTP"].drop_duplicates(["equipamento", "funcao"])
        for _, r in alvo.iterrows():
            cat.add("E17", r.to_dict(), None, ACAO_PEND, f"Ponto {p} usado por: {desc}.{extra}")

    # E18 funções de CDV sem nenhuma ligação
    for eq, g in df[df["tipo"] == "CDV"].groupby("equipamento"):
        for fn in FUNCOES_CDV:
            gf = g[g["funcao"] == fn]
            if fn == "N12-CTP":
                continue
            if not gf.empty and not (gf["valor"].notna() & gf["valor"].ne("-")).any():
                r = gf.iloc[0].to_dict(); r["valor_bruto"] = "'-' em todas as linhas"
                cat.add("E18", r, None, ACAO_PEND, f"{eq} não tem nenhum ponto para {fn}.")


# ---------------------------------------------------------------------------
# 3. Métricas do painel original (para comparação honesta)
# ---------------------------------------------------------------------------
def metricas_painel_original(bruto: pd.DataFrame):
    """Reproduz a lógica do painel anterior (A1.py): '-' vira vazio, sem validação."""
    b = bruto.copy()
    v = b["valor_bruto"].fillna("").str.replace(r"\s+", " ", regex=True).str.strip()
    v = v.where(~v.isin(["", "-", "—", "–"]))
    preench = v.notna()
    ids = v[preench]
    return dict(
        grupos_cdv=b.loc[(b["tipo"] == "CDV") & preench, "equipamento_bruto"].nunique(),
        grupos_sinal=b.loc[(b["tipo"] == "Sinal") & preench, "equipamento_bruto"].nunique(),
        preenchidos=int(preench.sum()),
        distintos=int(ids.nunique()),
        repetidos=int(ids.duplicated(keep=False).sum()),
        sem_id=int((~preench).sum()),
        posicoes=len(b),
    )


if __name__ == "__main__":
    r = processar()
    e = r["erros"]
    print(len(r["bruto"]), len(r["corrigido"]), len(e))
    print(e.groupby(["regra", "nome_regra", "acao"]).size())
    pd.set_option("display.width", 250, "display.max_colwidth", 90, "display.max_rows", 300)
    print(e[["id_erro", "regra", "equipamento", "celula", "funcao", "valor_bruto", "valor_corrigido", "acao", "detalhe"]].to_string())
    print(metricas_painel_original(r["bruto"]))
